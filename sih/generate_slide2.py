import sys
import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

def build_slide_2():
    prs = Presentation('sih/SIH-2026_template_backup.pptx')
    slide = prs.slides[1]
    
    # Identify persistent template shapes (Title, Picture 10, Footers)
    keep_shapes = []
    remove_shapes = []
    
    for s in slide.shapes:
        if s.name == 'Title 1' or s.name == 'Picture 10' or 'Placeholder' in s.name or s.name == 'Rectangle 8':
            keep_shapes.append(s)
        else:
            remove_shapes.append(s)
            
    for s in remove_shapes:
        sp = s._element
        sp.getparent().remove(sp)
        
    # Configure Slide Title
    if slide.shapes.title:
        title = slide.shapes.title
        title.text = "IDEA & PROPOSED SOLUTION"
        for p in title.text_frame.paragraphs:
            p.font.name = "Arial"
            p.font.size = Pt(24)
            p.font.bold = True
            p.font.color.rgb = RGBColor(0, 51, 102) # Navy #003366
            p.alignment = PP_ALIGN.LEFT
            
    # Slide is 13.333" x 7.5"
    # LEFT COLUMN: SIHex2 Style (Existing Problem, Proposed Solution, UVP)
    left_x = Inches(0.48)
    col_w = Inches(6.05)
    
    # -------------------------------------------------------------
    # 1. EXISTING PROBLEM
    # -------------------------------------------------------------
    pill1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, Inches(1.15), Inches(2.20), Inches(0.32))
    pill1.fill.solid()
    pill1.fill.fore_color.rgb = RGBColor(90, 107, 124) # Slate Blue #5A6B7C
    pill1.line.color.rgb = RGBColor(70, 85, 100)
    tf1 = pill1.text_frame
    tf1.word_wrap = False
    tf1.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf1.paragraphs[0]
    p.text = "PROBLEM EXISTING"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    
    box1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, Inches(1.50), col_w, Inches(1.50))
    box1.fill.solid()
    box1.fill.fore_color.rgb = RGBColor(248, 250, 252) # Soft clean off-white
    box1.line.color.rgb = RGBColor(203, 213, 225) # Slate border
    box1.line.width = Pt(1.2)
    tf_b1 = box1.text_frame
    tf_b1.word_wrap = True
    tf_b1.vertical_anchor = MSO_ANCHOR.TOP
    tf_b1.margin_left = Inches(0.18)
    tf_b1.margin_right = Inches(0.18)
    tf_b1.margin_top = Inches(0.12)
    tf_b1.margin_bottom = Inches(0.10)
    
    prob_items = [
        ("₹4,000+ Cr Escrow Lockup", "Sub-judice compensation under Section 3H stalls critical national corridors for 1,200+ days."),
        ("Section 3D Statutory Lapse", "Section 3A notifications automatically lapse at Day 365 if environmental/revenue clearances stall."),
        ("Disconnected Land Silos", "Gazette notices, e-Courts stays, cadastral Khasra maps, and civil road works operate in blind silos.")
    ]
    
    for i, (title_txt, desc_txt) in enumerate(prob_items):
        p = tf_b1.paragraphs[0] if i == 0 else tf_b1.add_paragraph()
        p.space_after = Pt(5)
        p.space_before = Pt(2)
        
        r_arrow = p.add_run()
        r_arrow.text = "➤  "
        r_arrow.font.name = "Arial"
        r_arrow.font.size = Pt(10)
        r_arrow.font.bold = True
        r_arrow.font.color.rgb = RGBColor(0, 112, 192) # SIHex2 blue
        
        r_title = p.add_run()
        r_title.text = f"{title_txt} : "
        r_title.font.name = "Arial"
        r_title.font.size = Pt(10)
        r_title.font.bold = True
        r_title.font.color.rgb = RGBColor(0, 112, 192)
        
        r_desc = p.add_run()
        r_desc.text = desc_txt
        r_desc.font.name = "Arial"
        r_desc.font.size = Pt(9.5)
        r_desc.font.color.rgb = RGBColor(30, 41, 59)
        
    # -------------------------------------------------------------
    # 2. PROPOSED SOLUTION
    # -------------------------------------------------------------
    pill2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, Inches(3.10), Inches(2.30), Inches(0.32))
    pill2.fill.solid()
    pill2.fill.fore_color.rgb = RGBColor(90, 107, 124)
    pill2.line.color.rgb = RGBColor(70, 85, 100)
    tf2 = pill2.text_frame
    tf2.word_wrap = False
    tf2.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf2.paragraphs[0]
    p.text = "PROPOSED SOLUTION"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    
    box2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, Inches(3.45), col_w, Inches(1.58))
    box2.fill.solid()
    box2.fill.fore_color.rgb = RGBColor(248, 250, 252)
    box2.line.color.rgb = RGBColor(203, 213, 225)
    box2.line.width = Pt(1.2)
    tf_b2 = box2.text_frame
    tf_b2.word_wrap = True
    tf_b2.vertical_anchor = MSO_ANCHOR.TOP
    tf_b2.margin_left = Inches(0.18)
    tf_b2.margin_right = Inches(0.18)
    tf_b2.margin_top = Inches(0.12)
    tf_b2.margin_bottom = Inches(0.10)
    
    sol_items = [
        ("Decoupled Central DB (<15ms)", "Pre-ingests 47 standardized statutory features across civil, revenue, and court databases."),
        ("Dual-Track Reality Engine", "Decouples physical civil road construction from court escrow disputes under Section 3D(2) vesting."),
        ("TreeSHAP Statutory XAI", "Quantifies exact additive day delays per bottleneck and automates pre-drafted DFA legal notices.")
    ]
    
    for i, (title_txt, desc_txt) in enumerate(sol_items):
        p = tf_b2.paragraphs[0] if i == 0 else tf_b2.add_paragraph()
        p.space_after = Pt(5)
        p.space_before = Pt(2)
        r_arrow = p.add_run()
        r_arrow.text = "➤  "
        r_arrow.font.name = "Arial"
        r_arrow.font.size = Pt(10)
        r_arrow.font.bold = True
        r_arrow.font.color.rgb = RGBColor(0, 112, 192)
        
        r_title = p.add_run()
        r_title.text = f"{title_txt} : "
        r_title.font.name = "Arial"
        r_title.font.size = Pt(10)
        r_title.font.bold = True
        r_title.font.color.rgb = RGBColor(0, 112, 192)
        
        r_desc = p.add_run()
        r_desc.text = desc_txt
        r_desc.font.name = "Arial"
        r_desc.font.size = Pt(9.5)
        r_desc.font.color.rgb = RGBColor(30, 41, 59)
        
    # -------------------------------------------------------------
    # 3. UVP (UNIQUE VALUE PROPOSITION)
    # -------------------------------------------------------------
    pill3 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, Inches(5.12), Inches(3.20), Inches(0.32))
    pill3.fill.solid()
    pill3.fill.fore_color.rgb = RGBColor(90, 107, 124)
    pill3.line.color.rgb = RGBColor(70, 85, 100)
    tf3 = pill3.text_frame
    tf3.word_wrap = False
    tf3.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf3.paragraphs[0]
    p.text = "UVP (UNIQUE VALUE PROPOSITION)"
    p.font.name = "Arial"
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    
    box3 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, Inches(5.47), col_w, Inches(1.55))
    box3.fill.solid()
    box3.fill.fore_color.rgb = RGBColor(248, 250, 252)
    box3.line.color.rgb = RGBColor(203, 213, 225)
    box3.line.width = Pt(1.2)
    tf_b3 = box3.text_frame
    tf_b3.word_wrap = True
    tf_b3.vertical_anchor = MSO_ANCHOR.TOP
    tf_b3.margin_left = Inches(0.18)
    tf_b3.margin_right = Inches(0.18)
    tf_b3.margin_top = Inches(0.12)
    tf_b3.margin_bottom = Inches(0.10)
    
    uvp_items = [
        ("Constitutional HITL (Arts. 77 & 166)", "Strict human-in-the-loop governance; zero auto-execution without Aadhaar DSC e-Sign."),
        ("340.8 km Interactive RoW GIS", "Hierarchical drill-down: Macro Alignment ➔ 8 Packages ➔ 156 Villages ➔ Khasra Plots."),
        ("SHA-256 Tamper-Proof Audit", "Cryptographic block hashing permanently seals all risk scores and SOP orders for CAG/CVC scrutiny.")
    ]
    
    for i, (title_txt, desc_txt) in enumerate(uvp_items):
        p = tf_b3.paragraphs[0] if i == 0 else tf_b3.add_paragraph()
        p.space_after = Pt(5)
        p.space_before = Pt(2)
        r_arrow = p.add_run()
        r_arrow.text = "➤  "
        r_arrow.font.name = "Arial"
        r_arrow.font.size = Pt(10)
        r_arrow.font.bold = True
        r_arrow.font.color.rgb = RGBColor(0, 112, 192)
        
        r_title = p.add_run()
        r_title.text = f"{title_txt} : "
        r_title.font.name = "Arial"
        r_title.font.size = Pt(10)
        r_title.font.bold = True
        r_title.font.color.rgb = RGBColor(0, 112, 192)
        
        r_desc = p.add_run()
        r_desc.text = desc_txt
        r_desc.font.name = "Arial"
        r_desc.font.size = Pt(9.5)
        r_desc.font.color.rgb = RGBColor(30, 41, 59)

    # -------------------------------------------------------------
    # RIGHT COLUMN: SIHex1 Style Flowchart + Innovation Cards
    # -------------------------------------------------------------
    right_x = Inches(6.70)
    right_w = Inches(6.20)
    
    # 1. Flowchart Container Box (SIHex1 style)
    flow_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_x, Inches(1.15), right_w, Inches(2.95))
    flow_box.fill.solid()
    flow_box.fill.fore_color.rgb = RGBColor(255, 255, 255)
    flow_box.line.color.rgb = RGBColor(203, 213, 225)
    flow_box.line.width = Pt(1.2)
    
    # Flowchart Title inside box
    f_title = slide.shapes.add_textbox(right_x, Inches(1.20), right_w, Inches(0.30))
    ft = f_title.text_frame
    p = ft.paragraphs[0]
    p.text = "GATISHAKTI AI End-to-end workflow"
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RGBColor(15, 23, 42)
    p.alignment = PP_ALIGN.CENTER
    
    # Flowchart Nodes (SIHex1 color coding: Pink start/end, Yellow/Gold process, Blue diamond decision)
    # Start node (Pink pill)
    n_start = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.55), Inches(1.05), Inches(0.40))
    n_start.fill.solid()
    n_start.fill.fore_color.rgb = RGBColor(248, 187, 208) # Pink #F8BBD0
    n_start.line.color.rgb = RGBColor(194, 24, 91)
    n_start.line.width = Pt(1)
    tf = n_start.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "Start Ingestion"
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(136, 14, 79)
    p.alignment = PP_ALIGN.CENTER
    
    # Arrow 1
    arr1 = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(7.98), Inches(1.68), Inches(0.22), Inches(0.14))
    arr1.fill.solid()
    arr1.fill.fore_color.rgb = RGBColor(100, 116, 139)
    arr1.line.color.rgb = RGBColor(100, 116, 139)
    
    # Node 1: Multi-Modal Ingestion (Yellow/Gold)
    n1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.25), Inches(1.52), Inches(1.85), Inches(0.46))
    n1.fill.solid()
    n1.fill.fore_color.rgb = RGBColor(255, 249, 196) # Light yellow
    n1.line.color.rgb = RGBColor(251, 192, 45)
    n1.line.width = Pt(1)
    tf = n1.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "Multi-Modal Ingestion\n(Gazette, e-Courts, Revenue)"
    p.font.size = Pt(8)
    p.font.bold = True
    p.font.color.rgb = RGBColor(93, 64, 55)
    p.alignment = PP_ALIGN.CENTER
    
    # Arrow 2
    arr2 = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(10.18), Inches(1.68), Inches(0.22), Inches(0.14))
    arr2.fill.solid()
    arr2.fill.fore_color.rgb = RGBColor(100, 116, 139)
    arr2.line.color.rgb = RGBColor(100, 116, 139)
    
    # Node 2: Decoupled Central DB (Gold)
    n2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.45), Inches(1.52), Inches(1.75), Inches(0.46))
    n2.fill.solid()
    n2.fill.fore_color.rgb = RGBColor(255, 224, 130) # Warm gold
    n2.line.color.rgb = RGBColor(255, 160, 0)
    n2.line.width = Pt(1)
    tf = n2.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "Decoupled Central DB\n(47 Standardized Features)"
    p.font.size = Pt(8)
    p.font.bold = True
    p.font.color.rgb = RGBColor(62, 39, 35)
    p.alignment = PP_ALIGN.CENTER
    
    # Arrow down from Node 2 to ML Core
    arr3 = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(11.25), Inches(2.02), Inches(0.14), Inches(0.18))
    arr3.fill.solid()
    arr3.fill.fore_color.rgb = RGBColor(100, 116, 139)
    arr3.line.color.rgb = RGBColor(100, 116, 139)
    
    # Row 2: ML Regressor -> Decision Diamond -> Branches
    n_ml = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.28), Inches(2.24), Inches(2.00), Inches(0.44))
    n_ml.fill.solid()
    n_ml.fill.fore_color.rgb = RGBColor(255, 224, 130)
    n_ml.line.color.rgb = RGBColor(255, 160, 0)
    n_ml.line.width = Pt(1)
    tf = n_ml.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "LightGBM + XGBoost Regressor\n(Predicts Residual Delay Days)"
    p.font.size = Pt(7.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(62, 39, 35)
    p.alignment = PP_ALIGN.CENTER
    
    # Arrow left from ML to Decision Diamond
    arr4 = slide.shapes.add_shape(MSO_SHAPE.LEFT_ARROW, Inches(10.02), Inches(2.38), Inches(0.20), Inches(0.14))
    arr4.fill.solid()
    arr4.fill.fore_color.rgb = RGBColor(100, 116, 139)
    arr4.line.color.rgb = RGBColor(100, 116, 139)
    
    # Decision Diamond (Cobalt Blue Diamond)
    n_dia = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(8.55), Inches(2.10), Inches(1.40), Inches(0.72))
    n_dia.fill.solid()
    n_dia.fill.fore_color.rgb = RGBColor(21, 101, 192) # SIHex1 Diamond Blue
    n_dia.line.color.rgb = RGBColor(13, 71, 161)
    tf = n_dia.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "Delay Risk\n> 60 Days?"
    p.font.size = Pt(7.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    
    # Branch 1 (Left: No -> Fast Track)
    lbl_no = slide.shapes.add_textbox(Inches(8.15), Inches(2.18), Inches(0.35), Inches(0.20))
    tf_no = lbl_no.text_frame
    p_no = tf_no.paragraphs[0]
    p_no.text = "No"
    p_no.font.size = Pt(7.5)
    p_no.font.bold = True
    p_no.font.color.rgb = RGBColor(46, 125, 50)
    
    arr_no = slide.shapes.add_shape(MSO_SHAPE.LEFT_ARROW, Inches(8.30), Inches(2.38), Inches(0.20), Inches(0.14))
    arr_no.fill.solid()
    arr_no.fill.fore_color.rgb = RGBColor(46, 125, 50) # Green arrow
    arr_no.line.color.rgb = RGBColor(46, 125, 50)
    
    n_no = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(2.24), Inches(1.40), Inches(0.44))
    n_no.fill.solid()
    n_no.fill.fore_color.rgb = RGBColor(232, 245, 233) # Light green
    n_no.line.color.rgb = RGBColor(76, 175, 80)
    n_no.line.width = Pt(1)
    tf = n_no.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "Fast-Track Notice\n(Sec 3A ➔ 3D)"
    p.font.size = Pt(7.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(27, 94, 32)
    p.alignment = PP_ALIGN.CENTER
    
    # Branch 2 (Down: Yes -> TreeSHAP Attribution & SOP)
    lbl_yes = slide.shapes.add_textbox(Inches(9.28), Inches(2.70), Inches(0.35), Inches(0.20))
    tf_yes = lbl_yes.text_frame
    p_yes = tf_yes.paragraphs[0]
    p_yes.text = "Yes"
    p_yes.font.size = Pt(7.5)
    p_yes.font.bold = True
    p_yes.font.color.rgb = RGBColor(230, 81, 0)
    
    arr_yes = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(9.18), Inches(2.86), Inches(0.14), Inches(0.18))
    arr_yes.fill.solid()
    arr_yes.fill.fore_color.rgb = RGBColor(230, 81, 0)
    arr_yes.line.color.rgb = RGBColor(230, 81, 0)
    
    # Row 3: TreeSHAP SOP -> Officer e-Sign -> Audit Sealed
    n_sop = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(3.08), Inches(1.88), Inches(0.46))
    n_sop.fill.solid()
    n_sop.fill.fore_color.rgb = RGBColor(255, 224, 178) # Light orange
    n_sop.line.color.rgb = RGBColor(251, 140, 0)
    n_sop.line.width = Pt(1)
    tf = n_sop.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "TreeSHAP Attribution\n+ Automated Legal DFA SOP"
    p.font.size = Pt(7.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(191, 54, 12)
    p.alignment = PP_ALIGN.CENTER
    
    arr5 = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(8.80), Inches(3.24), Inches(0.20), Inches(0.14))
    arr5.fill.solid()
    arr5.fill.fore_color.rgb = RGBColor(100, 116, 139)
    arr5.line.color.rgb = RGBColor(100, 116, 139)
    
    n_hitl = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.05), Inches(3.08), Inches(1.95), Inches(0.46))
    n_hitl.fill.solid()
    n_hitl.fill.fore_color.rgb = RGBColor(255, 236, 179) # Amber
    n_hitl.line.color.rgb = RGBColor(255, 179, 0)
    n_hitl.line.width = Pt(1)
    tf = n_hitl.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "Officer HITL Review\n+ Aadhaar DSC e-Sign"
    p.font.size = Pt(7.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(78, 52, 46)
    p.alignment = PP_ALIGN.CENTER
    
    arr6 = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(11.06), Inches(3.24), Inches(0.18), Inches(0.14))
    arr6.fill.solid()
    arr6.fill.fore_color.rgb = RGBColor(100, 116, 139)
    arr6.line.color.rgb = RGBColor(100, 116, 139)
    
    n_end = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.30), Inches(3.08), Inches(1.45), Inches(0.46))
    n_end.fill.solid()
    n_end.fill.fore_color.rgb = RGBColor(248, 187, 208) # Pink
    n_end.line.color.rgb = RGBColor(194, 24, 91)
    n_end.line.width = Pt(1)
    tf = n_end.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "CAG-Ready\nSHA-256 Ledger"
    p.font.size = Pt(7.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(136, 14, 79)
    p.alignment = PP_ALIGN.CENTER
    
    # -------------------------------------------------------------
    # 2. Innovation and Uniqueness Banner & 5 Vertical Cards (SIHex1 style)
    # -------------------------------------------------------------
    inno_pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.30), Inches(4.22), Inches(3.00), Inches(0.34))
    inno_pill.fill.solid()
    inno_pill.fill.fore_color.rgb = RGBColor(75, 111, 68) # Olive Green #4B6F44
    inno_pill.line.color.rgb = RGBColor(56, 85, 51)
    tf_ip = inno_pill.text_frame
    tf_ip.word_wrap = False
    tf_ip.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf_ip.paragraphs[0]
    p.text = "Innovation and Uniqueness"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    
    card_data = [
        {
            "title": "Decoupled DB",
            "tag": "<15ms Direct Query",
            "desc": "Direct normalized cache of 47 statutory attributes eliminates slow manual OCR bottlenecks.",
            "icon": "⚡ Fast Cache",
            "bg": RGBColor(168, 68, 60),      # Terracotta Red
            "border": RGBColor(235, 180, 175)
        },
        {
            "title": "Dual-Track ML",
            "tag": "Civil vs Legal",
            "desc": "Decouples active physical road works from sub-judice court disputes under Sec 3D(2) vesting.",
            "icon": "⚖️ Sec 3D(2)",
            "bg": RGBColor(166, 107, 36),     # Amber Brown
            "border": RGBColor(240, 205, 160)
        },
        {
            "title": "TreeSHAP XAI",
            "tag": "Statutory Attribution",
            "desc": "Quantifies exact additive day delays per bottleneck and automates pre-drafted legal DFA SOPs.",
            "icon": "🔍 Factor XAI",
            "bg": RGBColor(118, 126, 52),     # Olive Gold
            "border": RGBColor(215, 220, 170)
        },
        {
            "title": "340.8km GIS",
            "tag": "Corridor Telemetry",
            "desc": "Hierarchical drill-down: 8 Packages, 156 Villages, and Khasra plots with chainage telemetry.",
            "icon": "🗺️ RoW GIS",
            "bg": RGBColor(196, 85, 55),      # Coral Orange
            "border": RGBColor(245, 190, 175)
        },
        {
            "title": "Const. HITL",
            "tag": "Arts. 77 & 166",
            "desc": "Aadhaar DSC e-Sign validation with immutable SHA-256 cryptographic audit trails for CAG.",
            "icon": "🛡️ SHA-256",
            "bg": RGBColor(84, 110, 122),     # Slate Steel
            "border": RGBColor(190, 205, 215)
        }
    ]
    
    card_w = Inches(1.15)
    card_gap = Inches(0.11)
    card_top = Inches(4.62)
    card_h = Inches(2.40)
    
    for i, c in enumerate(card_data):
        cx = right_x + i * (card_w + card_gap)
        
        # Main Card shape
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, card_top, card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = c["bg"]
        card.line.color.rgb = c["border"]
        card.line.width = Pt(1.5)
        
        tf_c = card.text_frame
        tf_c.word_wrap = True
        tf_c.vertical_anchor = MSO_ANCHOR.TOP
        tf_c.margin_left = Inches(0.06)
        tf_c.margin_right = Inches(0.06)
        tf_c.margin_top = Inches(0.08)
        tf_c.margin_bottom = Inches(0.06)
        
        # Title
        p0 = tf_c.paragraphs[0]
        p0.text = c["title"]
        p0.font.name = "Arial"
        p0.font.size = Pt(8.5)
        p0.font.bold = True
        p0.font.color.rgb = RGBColor(255, 255, 255)
        p0.alignment = PP_ALIGN.CENTER
        p0.space_after = Pt(2)
        
        # Tag
        p_tag = tf_c.add_paragraph()
        p_tag.text = f"[{c['tag']}]"
        p_tag.font.name = "Arial"
        p_tag.font.size = Pt(7)
        p_tag.font.bold = True
        p_tag.font.color.rgb = RGBColor(255, 255, 200)
        p_tag.alignment = PP_ALIGN.CENTER
        p_tag.space_after = Pt(4)
        
        # Desc
        p_desc = tf_c.add_paragraph()
        p_desc.text = c["desc"]
        p_desc.font.name = "Arial"
        p_desc.font.size = Pt(7)
        p_desc.font.color.rgb = RGBColor(245, 245, 245)
        p_desc.alignment = PP_ALIGN.CENTER
        p_desc.space_after = Pt(4)
        
        # Bottom Badge shape inside card
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx + Inches(0.08), card_top + card_h - Inches(0.38), card_w - Inches(0.16), Inches(0.28))
        badge.fill.solid()
        badge.fill.fore_color.rgb = RGBColor(255, 255, 255)
        badge.line.color.rgb = c["border"]
        badge.line.width = Pt(0.8)
        tf_b = badge.text_frame
        tf_b.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_b = tf_b.paragraphs[0]
        p_b.text = c["icon"]
        p_b.font.name = "Arial"
        p_b.font.size = Pt(7)
        p_b.font.bold = True
        p_b.font.color.rgb = c["bg"]
        p_b.alignment = PP_ALIGN.CENTER

    prs.save('sih/SIH-2026_template_backup.pptx')
    prs.save('sih/GatiShakti_AI_SIH2026_Presentation.pptx')
    print("Slide 2 refined & saved successfully!")

if __name__ == '__main__':
    build_slide_2()
