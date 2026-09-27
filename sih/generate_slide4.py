import sys
import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

def build_slide_4():
    prs = Presentation('sih/SIH-2026_template_backup.pptx')
    slide = prs.slides[3]
    
    # Identify persistent template shapes (Title, Picture 10, Footers)
    keep_shapes = []
    remove_shapes = []
    
    for s in slide.shapes:
        if s.name == 'Title 1' or s.name == 'Picture 10' or 'Placeholder' in s.name or s.name == 'Rectangle 9':
            keep_shapes.append(s)
        else:
            remove_shapes.append(s)
            
    for s in remove_shapes:
        sp = s._element
        sp.getparent().remove(sp)
        
    # Configure Slide Title
    if slide.shapes.title:
        title = slide.shapes.title
        title.text = "FEASIBILITY AND VIABILITY"
        for p in title.text_frame.paragraphs:
            p.font.name = "Arial"
            p.font.size = Pt(24)
            p.font.bold = True
            p.font.color.rgb = RGBColor(0, 51, 102) # Navy #003366
            p.alignment = PP_ALIGN.LEFT
            
    # Dimensions: 13.333" x 7.5"
    # TOP HALF: Two Large Cards (Left: Feasibility, Right: Viability)
    top_y = Inches(1.15)
    top_h = Inches(3.40)
    col_w = Inches(6.08)
    left_x = Inches(0.48)
    right_x = Inches(6.78)
    
    # -------------------------------------------------------------
    # 1. TOP-LEFT CARD: FEASIBILITY
    # -------------------------------------------------------------
    card_feas = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, top_y, col_w, top_h)
    card_feas.fill.solid()
    card_feas.fill.fore_color.rgb = RGBColor(255, 255, 255)
    card_feas.line.color.rgb = RGBColor(0, 112, 192) # Vibrant SIHex2 Blue
    card_feas.line.width = Pt(1.5)
    
    # Header Icon & Text
    slide.shapes.add_picture('sih/assets/feasibility_badge.png', left_x + Inches(0.16), top_y + Inches(0.12), width=Inches(0.36), height=Inches(0.36))
    
    hdr_feas = slide.shapes.add_textbox(left_x + Inches(0.56), top_y + Inches(0.10), Inches(4.5), Inches(0.38))
    tf_hf = hdr_feas.text_frame
    p = tf_hf.paragraphs[0]
    p.text = "FEASIBILITY :"
    p.font.name = "Arial"
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 112, 192)
    
    # Content Box for Feasibility
    txt_feas = slide.shapes.add_textbox(left_x + Inches(0.16), top_y + Inches(0.50), col_w - Inches(0.32), top_h - Inches(0.55))
    tf_f = txt_feas.text_frame
    tf_f.word_wrap = True
    tf_f.vertical_anchor = MSO_ANCHOR.TOP
    tf_f.margin_top = Inches(0.02)
    tf_f.margin_left = Inches(0.04)
    tf_f.margin_right = Inches(0.04)
    
    feas_sections = [
        ("Technical :", [
            ("Proven Decoupled Stack : ", "Python 3.13, FastAPI, React 19, LightGBM, PostgreSQL — queries execute in <15ms."),
            ("Statutory Compliance : ", "100% compliant with National Highways Act 1956 (§ 3A–3H) and RFCTLARR Act 2013.")
        ]),
        ("Innovation :", [
            ("Dual-Track Reality Engine : ", "Decouples physical civil road construction from court escrow disputes under Sec 3D(2) vesting.")
        ]),
        ("Operational :", [
            ("Plug-and-Play Integration : ", "Point 11 compliant REST API Gateway integrates directly into existing NHAI / BhoomiRashi systems.")
        ]),
        ("Economic :", [
            ("Zero Recurring Licensing : ", "Built on 100% open-source software stack with no proprietary per-seat vendor fees.")
        ])
    ]
    
    first_p = True
    for sec_title, bullet_items in feas_sections:
        p_sec = tf_f.paragraphs[0] if first_p else tf_f.add_paragraph()
        first_p = False
        p_sec.text = sec_title
        p_sec.font.name = "Arial"
        p_sec.font.size = Pt(9.5)
        p_sec.font.bold = True
        p_sec.font.color.rgb = RGBColor(15, 23, 42)
        p_sec.space_before = Pt(3)
        p_sec.space_after = Pt(1)
        
        for b_label, b_desc in bullet_items:
            p_b = tf_f.add_paragraph()
            p_b.space_before = Pt(1)
            p_b.space_after = Pt(2)
            
            r_bullet = p_b.add_run()
            r_bullet.text = "•  "
            r_bullet.font.name = "Arial"
            r_bullet.font.size = Pt(8.5)
            r_bullet.font.bold = True
            r_bullet.font.color.rgb = RGBColor(0, 112, 192)
            
            r_lbl = p_b.add_run()
            r_lbl.text = b_label
            r_lbl.font.name = "Arial"
            r_lbl.font.size = Pt(8.5)
            r_lbl.font.bold = True
            r_lbl.font.color.rgb = RGBColor(0, 112, 192)
            
            r_txt = p_b.add_run()
            r_txt.text = b_desc
            r_txt.font.name = "Arial"
            r_txt.font.size = Pt(8.5)
            r_txt.font.color.rgb = RGBColor(51, 65, 85)

    # -------------------------------------------------------------
    # 2. TOP-RIGHT CARD: VIABILITY
    # -------------------------------------------------------------
    card_viab = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_x, top_y, col_w, top_h)
    card_viab.fill.solid()
    card_viab.fill.fore_color.rgb = RGBColor(255, 255, 255)
    card_viab.line.color.rgb = RGBColor(5, 150, 105) # Emerald Green
    card_viab.line.width = Pt(1.5)
    
    # Header Icon & Text
    slide.shapes.add_picture('sih/assets/viability_badge.png', right_x + Inches(0.16), top_y + Inches(0.12), width=Inches(0.36), height=Inches(0.36))
    
    hdr_viab = slide.shapes.add_textbox(right_x + Inches(0.56), top_y + Inches(0.10), Inches(4.5), Inches(0.38))
    tf_hv = hdr_viab.text_frame
    p = tf_hv.paragraphs[0]
    p.text = "VIABILITY :"
    p.font.name = "Arial"
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(5, 150, 105)
    
    # Content Box for Viability
    txt_viab = slide.shapes.add_textbox(right_x + Inches(0.16), top_y + Inches(0.50), col_w - Inches(0.32), top_h - Inches(0.55))
    tf_v = txt_viab.text_frame
    tf_v.word_wrap = True
    tf_v.vertical_anchor = MSO_ANCHOR.TOP
    tf_v.margin_top = Inches(0.02)
    tf_v.margin_left = Inches(0.04)
    tf_v.margin_right = Inches(0.04)
    
    viab_sections = [
        ("National Market & Scale :", 
         "₹4,000+ Cr in stalled Section 3H escrow funds across 1,200+ km of national highway corridors; directly unfreezes national CAPEX."),
        ("Policy & Statutory Alignment :", 
         "Perfectly synergizes with PM GatiShakti National Master Plan, Cabinet Secretariat guidelines, and digital land modernization."),
        ("Investment & National ROI :", 
         "Eliminates contractor idling compensation claims (saving ₹4.5 Cr per month per delayed package); instantaneous fiscal breakeven."),
        ("Scalability & Interstate Adoption :", 
         "Scales horizontally across all 28 States and 8 Union Territories via universal state-land ontological harmonizer.")
    ]
    
    first_p = True
    for sec_title, sec_desc in viab_sections:
        p_sec = tf_v.paragraphs[0] if first_p else tf_v.add_paragraph()
        first_p = False
        p_sec.space_before = Pt(3)
        p_sec.space_after = Pt(2)
        
        r_lbl = p_sec.add_run()
        r_lbl.text = f"{sec_title}  "
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(9.5)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = RGBColor(15, 23, 42)
        
        r_txt = p_sec.add_run()
        r_txt.text = sec_desc
        r_txt.font.name = "Arial"
        r_txt.font.size = Pt(8.5)
        r_txt.font.color.rgb = RGBColor(51, 65, 85)

    # =============================================================
    # 3. BOTTOM HALF: CHALLENGES & MITIGATION CARDS (SIHex2 style)
    # =============================================================
    bot_y = Inches(4.70)
    bot_h = Inches(1.98)
    
    # -------------------------------------------------------------
    # BOTTOM-LEFT: TECHNICAL CHALLENGES
    # -------------------------------------------------------------
    card_tch = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, bot_y, col_w, bot_h)
    card_tch.fill.solid()
    card_tch.fill.fore_color.rgb = RGBColor(255, 255, 255)
    card_tch.line.color.rgb = RGBColor(0, 112, 192)
    card_tch.line.width = Pt(1.5)
    
    slide.shapes.add_picture('sih/assets/challenge_badge.png', left_x + Inches(0.16), bot_y + Inches(0.12), width=Inches(0.32), height=Inches(0.32))
    
    hdr_tch = slide.shapes.add_textbox(left_x + Inches(0.54), bot_y + Inches(0.10), Inches(5.0), Inches(0.35))
    tf_ht = hdr_tch.text_frame
    p = tf_ht.paragraphs[0]
    p.text = "TECHNICAL CHALLENGES :"
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 112, 192)
    
    txt_tch = slide.shapes.add_textbox(left_x + Inches(0.16), bot_y + Inches(0.48), col_w - Inches(0.32), bot_h - Inches(0.52))
    tf_tc = txt_tch.text_frame
    tf_tc.word_wrap = True
    tf_tc.vertical_anchor = MSO_ANCHOR.TOP
    tf_tc.margin_left = Inches(0.04)
    tf_tc.margin_right = Inches(0.04)
    tf_tc.margin_top = Inches(0.02)
    
    p_risk1 = tf_tc.paragraphs[0]
    p_risk1.space_after = Pt(4)
    r1 = p_risk1.add_run()
    r1.text = "Risk : "
    r1.font.bold = True
    r1.font.size = Pt(9)
    r1.font.color.rgb = RGBColor(15, 23, 42)
    r2 = p_risk1.add_run()
    r2.text = "Heterogeneous state land nomenclature (Khasra, Khatauni, Jamabandi) and missing cadastral coordinate data."
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = RGBColor(51, 65, 85)
    
    p_mit1 = tf_tc.add_paragraph()
    p_mit1.space_before = Pt(2)
    m1 = p_mit1.add_run()
    m1.text = "Mitigation : "
    m1.font.bold = True
    m1.font.size = Pt(9)
    m1.font.color.rgb = RGBColor(0, 112, 192)
    m2 = p_mit1.add_run()
    m2.text = "Universal Ontological Harmonizer normalizes all 47 state parameters into unified schema with automated missing-value imputation."
    m2.font.size = Pt(8.5)
    m2.font.color.rgb = RGBColor(51, 65, 85)

    # -------------------------------------------------------------
    # BOTTOM-RIGHT: BUSINESS / STATUTORY CHALLENGES
    # -------------------------------------------------------------
    card_bch = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_x, bot_y, col_w, bot_h)
    card_bch.fill.solid()
    card_bch.fill.fore_color.rgb = RGBColor(255, 255, 255)
    card_bch.line.color.rgb = RGBColor(5, 150, 105)
    card_bch.line.width = Pt(1.5)
    
    slide.shapes.add_picture('sih/assets/governance_badge.png', right_x + Inches(0.16), bot_y + Inches(0.12), width=Inches(0.32), height=Inches(0.32))
    
    hdr_bch = slide.shapes.add_textbox(right_x + Inches(0.54), bot_y + Inches(0.10), Inches(5.0), Inches(0.35))
    tf_hb = hdr_bch.text_frame
    p = tf_hb.paragraphs[0]
    p.text = "STATUTORY & ADOPTION CHALLENGES :"
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(5, 150, 105)
    
    txt_bch = slide.shapes.add_textbox(right_x + Inches(0.16), bot_y + Inches(0.48), col_w - Inches(0.32), bot_h - Inches(0.52))
    tf_bc = txt_bch.text_frame
    tf_bc.word_wrap = True
    tf_bc.vertical_anchor = MSO_ANCHOR.TOP
    tf_bc.margin_left = Inches(0.04)
    tf_bc.margin_right = Inches(0.04)
    tf_bc.margin_top = Inches(0.02)
    
    p_risk2 = tf_bc.paragraphs[0]
    p_risk2.space_after = Pt(4)
    r1 = p_risk2.add_run()
    r1.text = "Risk : "
    r1.font.bold = True
    r1.font.size = Pt(9)
    r1.font.color.rgb = RGBColor(15, 23, 42)
    r2 = p_risk2.add_run()
    r2.text = "Bureaucratic apprehension towards automated 'black-box' decisions and pending Section 3H court litigation stays."
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = RGBColor(51, 65, 85)
    
    p_mit2 = tf_bc.add_paragraph()
    p_mit2.space_before = Pt(2)
    m1 = p_mit2.add_run()
    m1.text = "Mitigation : "
    m1.font.bold = True
    m1.font.size = Pt(9)
    m1.font.color.rgb = RGBColor(5, 150, 105)
    m2 = p_mit2.add_run()
    m2.text = "Constitutional HITL preserves sovereign statutory authority under Articles 77 & 166 with Aadhaar DSC e-Sign, backed by TreeSHAP explainability."
    m2.font.size = Pt(8.5)
    m2.font.color.rgb = RGBColor(51, 65, 85)

    prs.save('sih/SIH-2026_template_backup.pptx')
    prs.save('sih/GatiShakti_AI_SIH2026_Presentation.pptx')
    print("Slide 4 built successfully in both PPTX files!")

if __name__ == '__main__':
    build_slide_4()
