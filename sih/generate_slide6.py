import sys
import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

def build_slide_6():
    prs = Presentation('sih/SIH-2026_template_backup.pptx')
    slide = prs.slides[5]
    
    # Identify persistent template shapes (Title, Picture 11, Footers)
    keep_shapes = []
    remove_shapes = []
    
    for s in slide.shapes:
        if s.name == 'Title 1' or s.name == 'Picture 11' or 'Placeholder' in s.name or s.name == 'Rectangle 9':
            keep_shapes.append(s)
        else:
            remove_shapes.append(s)
            
    for s in remove_shapes:
        sp = s._element
        sp.getparent().remove(sp)
        
    # Configure Slide Title
    if slide.shapes.title:
        title = slide.shapes.title
        title.text = "RESEARCH AND REFERENCES"
        for p in title.text_frame.paragraphs:
            p.font.name = "Arial"
            p.font.size = Pt(24)
            p.font.bold = True
            p.font.color.rgb = RGBColor(0, 51, 102) # Navy #003366
            p.alignment = PP_ALIGN.LEFT
            
    # Slide dimensions: 13.333" x 7.5"
    left_x = Inches(0.48)
    full_w = Inches(12.36)
    
    # =============================================================
    # 1. TOP CARD: STATUTORY & RESEARCH CITATIONS (SIHex2 style)
    # =============================================================
    top_y = Inches(1.15)
    top_h = Inches(1.50)
    
    card_top = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, top_y, full_w, top_h)
    card_top.fill.solid()
    card_top.fill.fore_color.rgb = RGBColor(255, 255, 255)
    card_top.line.color.rgb = RGBColor(203, 213, 225)
    card_top.line.width = Pt(1.2)
    
    # Add Book / Law Icon
    slide.shapes.add_picture('sih/assets/statutory_book_icon.png', left_x + Inches(0.16), top_y + Inches(0.12), width=Inches(0.42), height=Inches(0.42))
    
    # Citations Text Box
    txt_top = slide.shapes.add_textbox(left_x + Inches(0.68), top_y + Inches(0.06), full_w - Inches(0.80), top_h - Inches(0.12))
    tf_top = txt_top.text_frame
    tf_top.word_wrap = True
    tf_top.vertical_anchor = MSO_ANCHOR.TOP
    tf_top.margin_left = Inches(0.02)
    tf_top.margin_right = Inches(0.04)
    tf_top.margin_top = Inches(0.02)
    
    citations = [
        ("The National Highways Act, 1956 & RFCTLARR Act, 2013 : ",
         "Foundational Indian statutory frameworks governing land acquisition (§ 3A to 3H) and mandatory 365-day Section 3D lapse timelines. Full Details : NH Act 1956 / RFCTLARR 2013 / PM GatiShakti Framework"),
        ("Explainable AI (XAI) & TreeSHAP : ",
         "Lundberg & Lee (NeurIPS 2017) foundational game-theoretic feature attribution; Chen & Guestrin (KDD 2016) XGBoost gradient boosted trees. Links : NeurIPS 2017 / XGBoost KDD 2016"),
        ("Constitutional Governance & Audit Rigor : ",
         "Adherence to Articles 77 & 166 of the Constitution of India; CVC & CAG compliance guidelines for digital infrastructure procurement. Links : MoRTH Directives / CAG Guidelines")
    ]
    
    for i, (c_label, c_desc) in enumerate(citations):
        p = tf_top.paragraphs[0] if i == 0 else tf_top.add_paragraph()
        p.space_after = Pt(2.5)
        p.space_before = Pt(1.5)
        
        r_b = p.add_run()
        r_b.text = "•  "
        r_b.font.name = "Arial"
        r_b.font.size = Pt(8.5)
        r_b.font.bold = True
        r_b.font.color.rgb = RGBColor(0, 112, 192)
        
        r_l = p.add_run()
        r_l.text = c_label
        r_l.font.name = "Arial"
        r_l.font.size = Pt(8.5)
        r_l.font.bold = True
        r_l.font.color.rgb = RGBColor(15, 23, 42)
        
        r_d = p.add_run()
        r_d.text = c_desc
        r_d.font.name = "Arial"
        r_d.font.size = Pt(8)
        r_d.font.color.rgb = RGBColor(51, 65, 85)

    # =============================================================
    # 2. MIDDLE SECTION: SOLUTIONS EXIST vs STANDS OUT (SIHex2 style)
    # =============================================================
    mid_y = Inches(2.78)
    mid_h = Inches(1.60)
    box_w1 = Inches(5.65)
    box_w2 = Inches(6.05)
    
    # Left Card: SOLUTIONS ALREADY EXIST
    card_exist = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, mid_y, box_w1, mid_h)
    card_exist.fill.solid()
    card_exist.fill.fore_color.rgb = RGBColor(240, 247, 255) # Soft Blue tint
    card_exist.line.color.rgb = RGBColor(187, 222, 251)
    card_exist.line.width = Pt(1.2)
    
    hdr_ex = slide.shapes.add_textbox(left_x, mid_y + Inches(0.08), box_w1, Inches(0.28))
    tf_hex = hdr_ex.text_frame
    p = tf_hex.paragraphs[0]
    p.text = "SOLUTIONS ALREADY EXIST"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 112, 192)
    p.alignment = PP_ALIGN.CENTER
    
    txt_ex = slide.shapes.add_textbox(left_x + Inches(0.16), mid_y + Inches(0.38), box_w1 - Inches(0.32), mid_h - Inches(0.44))
    tf_e = txt_ex.text_frame
    tf_e.word_wrap = True
    tf_e.vertical_anchor = MSO_ANCHOR.TOP
    tf_e.margin_left = Inches(0.04)
    tf_e.margin_right = Inches(0.04)
    
    p0 = tf_e.paragraphs[0]
    p0.space_after = Pt(4)
    r1 = p0.add_run()
    r1.text = "•  Most are fragmented, physical paper-based registers, isolated cadastral maps, or slow manual OCR tools built with zero statutory delay prediction.\n"
    r1.font.name = "Arial"
    r1.font.size = Pt(8.5)
    r1.font.color.rgb = RGBColor(51, 65, 85)
    
    p1 = tf_e.add_paragraph()
    r2 = p1.add_run()
    r2.text = "•  EXISTING SYSTEMS :  "
    r2.font.name = "Arial"
    r2.font.size = Pt(8.5)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(15, 23, 42)
    
    r3 = p1.add_run()
    r3.text = "BhoomiRashi  /  e-Courts WSDL  /  State Bhulekh  /  Tarang"
    r3.font.name = "Arial"
    r3.font.size = Pt(8.5)
    r3.font.color.rgb = RGBColor(0, 112, 192)

    # Center Transitional Right Arrow (SIHex2 style)
    arr_mid = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left_x + box_w1 + Inches(0.12), mid_y + Inches(0.60), Inches(0.42), Inches(0.32))
    arr_mid.fill.solid()
    arr_mid.fill.fore_color.rgb = RGBColor(15, 23, 42) # Solid Dark Navy/Black
    arr_mid.line.fill.background()

    # Right Card: OUR SOLUTION STANDS OUT
    card_ours = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x + box_w1 + Inches(0.66), mid_y, box_w2, mid_h)
    card_ours.fill.solid()
    card_ours.fill.fore_color.rgb = RGBColor(240, 247, 255)
    card_ours.line.color.rgb = RGBColor(187, 222, 251)
    card_ours.line.width = Pt(1.2)
    
    hdr_ou = slide.shapes.add_textbox(left_x + box_w1 + Inches(0.66), mid_y + Inches(0.08), box_w2, Inches(0.28))
    tf_hou = hdr_ou.text_frame
    p = tf_hou.paragraphs[0]
    p.text = "OUR SOLUTION STANDS OUT"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 112, 192)
    p.alignment = PP_ALIGN.CENTER
    
    txt_ou = slide.shapes.add_textbox(left_x + box_w1 + Inches(0.80), mid_y + Inches(0.38), box_w2 - Inches(0.28), mid_h - Inches(0.44))
    tf_o = txt_ou.text_frame
    tf_o.word_wrap = True
    tf_o.vertical_anchor = MSO_ANCHOR.TOP
    tf_o.margin_left = Inches(0.04)
    tf_o.margin_right = Inches(0.04)
    
    our_points = [
        ("Decoupled Central DB : ", "<15ms normalized query cache across 47 standardized attributes."),
        ("Dual-Track Reality Engine : ", "Separates active physical road works from court escrow disputes."),
        ("TreeSHAP Statutory XAI : ", "Quantifies exact additive day delays & auto-generates legal DFA SOP notices."),
        ("Constitutional HITL & Audit : ", "Aadhaar DSC e-Sign validation + SHA-256 tamper-proof ledger for CAG.")
    ]
    
    for i, (p_title, p_desc) in enumerate(our_points):
        p_pt = tf_o.paragraphs[0] if i == 0 else tf_o.add_paragraph()
        p_pt.space_after = Pt(2)
        
        r_b = p_pt.add_run()
        r_b.text = "•  "
        r_b.font.name = "Arial"
        r_b.font.size = Pt(8)
        r_b.font.bold = True
        r_b.font.color.rgb = RGBColor(0, 112, 192)
        
        r_t = p_pt.add_run()
        r_t.text = p_title
        r_t.font.name = "Arial"
        r_t.font.size = Pt(8)
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(15, 23, 42)
        
        r_d = p_pt.add_run()
        r_d.text = p_desc
        r_d.font.name = "Arial"
        r_d.font.size = Pt(8)
        r_d.font.color.rgb = RGBColor(51, 65, 85)

    # =============================================================
    # 3. BOTTOM CARD: PROJECT RESOURCES & REAL UI SCREENSHOTS
    # =============================================================
    bot_y = Inches(4.50)
    bot_h = Inches(2.25)
    
    card_bot = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, bot_y, full_w, bot_h)
    card_bot.fill.solid()
    card_bot.fill.fore_color.rgb = RGBColor(255, 255, 255)
    card_bot.line.color.rgb = RGBColor(203, 213, 225)
    card_bot.line.width = Pt(1.2)
    
    # Left Column: Project Links & Endpoints (w = 4.35")
    hdr_res = slide.shapes.add_textbox(left_x + Inches(0.16), bot_y + Inches(0.08), Inches(4.20), Inches(0.32))
    tf_hr = hdr_res.text_frame
    p = tf_hr.paragraphs[0]
    p.text = "PROJECT RESOURCES"
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 112, 192)
    
    txt_res = slide.shapes.add_textbox(left_x + Inches(0.16), bot_y + Inches(0.42), Inches(4.30), bot_h - Inches(0.50))
    tf_r = txt_res.text_frame
    tf_r.word_wrap = True
    tf_r.vertical_anchor = MSO_ANCHOR.TOP
    tf_r.margin_left = Inches(0.02)
    tf_r.margin_right = Inches(0.04)
    tf_r.margin_top = Inches(0.02)
    
    res_links = [
        ("GitHub Repository : ", "github.com/AryanSahu321/Land_Aquisition_Delay_pridictor"),
        ("REST API Gateway : ", "Live OpenAPI 3.0 swagger endpoints for government integration"),
        ("Interactive RoW GIS : ", "Full Purvanchal 340.8 km corridor with Khasra cadastral plots"),
        ("Audit Verification : ", "SHA-256 cryptographic block integrity ledger for CAG review")
    ]
    
    for i, (l_title, l_desc) in enumerate(res_links):
        p_l = tf_r.paragraphs[0] if i == 0 else tf_r.add_paragraph()
        p_l.space_after = Pt(4)
        p_l.space_before = Pt(1)
        
        r_t = p_l.add_run()
        r_t.text = l_title
        r_t.font.name = "Arial"
        r_t.font.size = Pt(8.5)
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(15, 23, 42)
        
        r_d = p_l.add_run()
        r_d.text = l_desc
        r_d.font.name = "Arial"
        r_d.font.size = Pt(8)
        r_d.font.color.rgb = RGBColor(0, 112, 192)

    # Right Column: Two Real Application UI Screenshots side-by-side!
    # Screen 1: Purvanchal GIS Corridor Map UI Preview
    screen1_x = left_x + Inches(4.55)
    screen_w1 = Inches(3.70)
    screen_h = Inches(2.05)
    slide.shapes.add_picture('sih/assets/ui_gis_corridor.png', screen1_x, bot_y + Inches(0.10), width=screen_w1, height=screen_h)
    
    # Screen 2: Predictive Delay & TreeSHAP Dashboard UI Preview
    screen2_x = screen1_x + screen_w1 + Inches(0.15)
    screen_w2 = Inches(3.85)
    slide.shapes.add_picture('sih/assets/ui_predictive_dashboard.png', screen2_x, bot_y + Inches(0.10), width=screen_w2, height=screen_h)

    prs.save('sih/SIH-2026_template_backup.pptx')
    prs.save('sih/GatiShakti_AI_SIH2026_Presentation.pptx')
    print("Slide 6 built successfully in both PPTX files!")

if __name__ == '__main__':
    build_slide_6()
