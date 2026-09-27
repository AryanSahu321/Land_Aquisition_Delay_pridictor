import sys
import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

def build_slide_3():
    prs = Presentation('sih/SIH-2026_template_backup.pptx')
    slide = prs.slides[2]
    
    # Identify persistent template shapes (Title, Picture 11, Footers)
    keep_shapes = []
    remove_shapes = []
    
    for s in slide.shapes:
        if s.name == 'Title 1' or s.name == 'Picture 11' or 'Placeholder' in s.name or s.name == 'Rectangle 9':
            keep_shapes.append(s)
        else:
            remove_shapes.append(s)
            
    for s in remove_shapes:
        sp = s._element
        sp.getparent().remove(sp)
        
    # Configure Slide Title
    if slide.shapes.title:
        title = slide.shapes.title
        title.text = "TECHNICAL APPROACH"
        for p in title.text_frame.paragraphs:
            p.font.name = "Arial"
            p.font.size = Pt(24)
            p.font.bold = True
            p.font.color.rgb = RGBColor(0, 51, 102) # Navy #003366
            p.alignment = PP_ALIGN.LEFT
            
    # Slide dimensions: 13.333" x 7.5"
    # LEFT & CENTER (Zones 1-4 + Tech Stack Bar): x = 0.45", w = 9.45"
    # RIGHT (Implementation Process): x = 10.05", w = 2.80"
    
    # =============================================================
    # 1. ARCHITECTURE ZONES (ZONES 1 - 4)
    # =============================================================
    arch_top = Inches(1.12)
    arch_h = Inches(4.55)
    
    # Outer container for architecture zones
    arch_bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.45), arch_top, Inches(9.45), arch_h)
    arch_bg.fill.solid()
    arch_bg.fill.fore_color.rgb = RGBColor(255, 255, 255)
    arch_bg.line.color.rgb = RGBColor(203, 213, 225)
    arch_bg.line.width = Pt(1.2)
    
    # -------------------------------------------------------------
    # ZONE 1: Users & External Services
    # -------------------------------------------------------------
    z1_x = Inches(0.50)
    z1_w = Inches(2.05)
    
    # Zone 1 Header
    z1_hdr = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, z1_x, arch_top + Inches(0.08), z1_w, Inches(0.40))
    z1_hdr.fill.solid()
    z1_hdr.fill.fore_color.rgb = RGBColor(227, 242, 253) # Light Blue
    z1_hdr.line.color.rgb = RGBColor(144, 202, 249)
    tf1 = z1_hdr.text_frame
    tf1.word_wrap = True
    tf1.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf1.paragraphs[0]
    p.text = "Zone 1: Users & External"
    p.font.name = "Arial"
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(13, 71, 161)
    p.alignment = PP_ALIGN.CENTER
    
    # Top Box: External Services
    z1_ext = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, z1_x, arch_top + Inches(0.56), z1_w, Inches(1.85))
    z1_ext.fill.solid()
    z1_ext.fill.fore_color.rgb = RGBColor(240, 247, 255)
    z1_ext.line.color.rgb = RGBColor(187, 222, 251)
    
    # Add Cloud Icon Image
    slide.shapes.add_picture('sih/assets/cloud_portals.png', z1_x + Inches(0.12), arch_top + Inches(0.62), width=Inches(0.48))
    
    # Text in External Services box
    tf_ext = z1_ext.text_frame
    tf_ext.word_wrap = True
    tf_ext.vertical_anchor = MSO_ANCHOR.TOP
    tf_ext.margin_top = Inches(0.06)
    tf_ext.margin_left = Inches(0.64) # Offset for icon
    tf_ext.margin_right = Inches(0.06)
    
    p0 = tf_ext.paragraphs[0]
    p0.text = "External Portals"
    p0.font.name = "Arial"
    p0.font.size = Pt(8.5)
    p0.font.bold = True
    p0.font.color.rgb = RGBColor(25, 118, 210)
    p0.space_after = Pt(2)
    
    ext_items = [
        "e-Courts WSDL API",
        "BhoomiRashi MoRTH",
        "State Jamabandi GIS",
        "PM GatiShakti NMP"
    ]
    for it in ext_items:
        p = tf_ext.add_paragraph()
        p.text = f"• {it}"
        p.font.name = "Arial"
        p.font.size = Pt(7.2)
        p.font.color.rgb = RGBColor(51, 65, 85)
        p.space_after = Pt(1.5)
        
    # Bottom Box: Administrative Roles (RBAC)
    z1_usr = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, z1_x, arch_top + Inches(2.50), z1_w, Inches(1.92))
    z1_usr.fill.solid()
    z1_usr.fill.fore_color.rgb = RGBColor(240, 247, 255)
    z1_usr.line.color.rgb = RGBColor(187, 222, 251)
    
    # Add Officer Avatar Image
    slide.shapes.add_picture('sih/assets/officer_avatar.png', z1_x + Inches(0.12), arch_top + Inches(2.58), width=Inches(0.46))
    
    tf_usr = z1_usr.text_frame
    tf_usr.word_wrap = True
    tf_usr.vertical_anchor = MSO_ANCHOR.TOP
    tf_usr.margin_top = Inches(0.06)
    tf_usr.margin_left = Inches(0.64)
    tf_usr.margin_right = Inches(0.06)
    
    p0 = tf_usr.paragraphs[0]
    p0.text = "Statutory Roles"
    p0.font.name = "Arial"
    p0.font.size = Pt(8.5)
    p0.font.bold = True
    p0.font.color.rgb = RGBColor(25, 118, 210)
    p0.space_after = Pt(2)
    
    usr_items = [
        "MoRTH & NPG Lead",
        "NHAI Project Director",
        "CALA (Land Officer)",
        "Field Tehsildars",
        "Landowner Farmers"
    ]
    for it in usr_items:
        p = tf_usr.add_paragraph()
        p.text = f"• {it}"
        p.font.name = "Arial"
        p.font.size = Pt(7.2)
        p.font.color.rgb = RGBColor(51, 65, 85)
        p.space_after = Pt(1.5)

    # Arrow 1: To Zone 2
    arr_z1 = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(2.58), arch_top + Inches(2.45), Inches(0.18), Inches(0.14))
    arr_z1.fill.solid()
    arr_z1.fill.fore_color.rgb = RGBColor(100, 116, 139)
    arr_z1.line.color.rgb = RGBColor(100, 116, 139)

    # -------------------------------------------------------------
    # ZONE 2: Frontend (Presentation Layer)
    # -------------------------------------------------------------
    z2_x = Inches(2.80)
    z2_w = Inches(2.15)
    
    z2_hdr = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, z2_x, arch_top + Inches(0.08), z2_w, Inches(0.40))
    z2_hdr.fill.solid()
    z2_hdr.fill.fore_color.rgb = RGBColor(232, 245, 233) # Mint green
    z2_hdr.line.color.rgb = RGBColor(165, 214, 167)
    tf2 = z2_hdr.text_frame
    tf2.word_wrap = True
    tf2.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf2.paragraphs[0]
    p.text = "Zone 2: Frontend (React 19)"
    p.font.name = "Arial"
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(27, 94, 32)
    p.alignment = PP_ALIGN.CENTER
    
    z2_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, z2_x, arch_top + Inches(0.56), z2_w, Inches(3.86))
    z2_box.fill.solid()
    z2_box.fill.fore_color.rgb = RGBColor(241, 248, 242)
    z2_box.line.color.rgb = RGBColor(165, 214, 167)
    
    # 5 Sub-component cards
    z2_cards = [
        ("Predictive Dashboard", "Real-time residual delay days, risk gauges & project cards."),
        ("340.8 km RoW GIS", "Leaflet GeoJSON EPSG:4326 drilldown across 8 packages & plots."),
        ("TreeSHAP Waterfall", "Visual factor attributions explaining statutory delay causes."),
        ("Prescriptive DFA SOP", "Automated legal notice drafts for Sec 3D/3G fast-tracking."),
        ("Constitutional HITL", "4-tier RBAC + Aadhaar DSC e-Sign gateway (Arts. 77/166).")
    ]
    
    card_h = Inches(0.68)
    for i, (title_c, desc_c) in enumerate(z2_cards):
        cy = arch_top + Inches(0.62) + i * Inches(0.75)
        c_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, z2_x + Inches(0.06), cy, z2_w - Inches(0.12), card_h)
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
        c_shape.line.color.rgb = RGBColor(200, 230, 201)
        tf_c = c_shape.text_frame
        tf_c.word_wrap = True
        tf_c.vertical_anchor = MSO_ANCHOR.TOP
        tf_c.margin_top = Inches(0.04)
        tf_c.margin_left = Inches(0.08)
        tf_c.margin_right = Inches(0.08)
        
        p0 = tf_c.paragraphs[0]
        p0.text = title_c
        p0.font.name = "Arial"
        p0.font.size = Pt(7.5)
        p0.font.bold = True
        p0.font.color.rgb = RGBColor(46, 125, 50)
        p0.space_after = Pt(1)
        
        p1 = tf_c.add_paragraph()
        p1.text = desc_c
        p1.font.name = "Arial"
        p1.font.size = Pt(6.5)
        p1.font.color.rgb = RGBColor(71, 85, 105)

    # Arrow 2: HTTP / REST
    arr_z2 = slide.shapes.add_shape(MSO_SHAPE.LEFT_RIGHT_ARROW, Inches(4.98), arch_top + Inches(2.45), Inches(0.20), Inches(0.14))
    arr_z2.fill.solid()
    arr_z2.fill.fore_color.rgb = RGBColor(100, 116, 139)
    arr_z2.line.color.rgb = RGBColor(100, 116, 139)

    # -------------------------------------------------------------
    # ZONE 3: Backend (Application Layer)
    # -------------------------------------------------------------
    z3_x = Inches(5.22)
    z3_w = Inches(2.20)
    
    z3_hdr = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, z3_x, arch_top + Inches(0.08), z3_w, Inches(0.40))
    z3_hdr.fill.solid()
    z3_hdr.fill.fore_color.rgb = RGBColor(237, 231, 246) # Lavender
    z3_hdr.line.color.rgb = RGBColor(179, 157, 219)
    tf3 = z3_hdr.text_frame
    tf3.word_wrap = True
    tf3.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf3.paragraphs[0]
    p.text = "Zone 3: Backend (FastAPI)"
    p.font.name = "Arial"
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(69, 39, 160)
    p.alignment = PP_ALIGN.CENTER
    
    z3_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, z3_x, arch_top + Inches(0.56), z3_w, Inches(3.86))
    z3_box.fill.solid()
    z3_box.fill.fore_color.rgb = RGBColor(246, 243, 251)
    z3_box.line.color.rgb = RGBColor(179, 157, 219)
    
    z3_cards = [
        ("REST API Gateway", "Point 11 compliant: async request routing, OpenAPI 3.0 & auth."),
        ("Decoupled Central DB", "<15ms cached responses for 47 standardized statutory features."),
        ("Dual-Track Engine", "Decouples physical civil road works from legal escrow disputes."),
        ("Statutory DFA Engine", "Generates legally binding Section 3A, 3D, 3G, 3H notice drafts."),
        ("SHA-256 Ledger Service", "Cryptographic block hashing ensuring immutable CAG audit trail.")
    ]
    
    for i, (title_c, desc_c) in enumerate(z3_cards):
        cy = arch_top + Inches(0.62) + i * Inches(0.75)
        c_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, z3_x + Inches(0.06), cy, z3_w - Inches(0.12), card_h)
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
        c_shape.line.color.rgb = RGBColor(209, 196, 233)
        tf_c = c_shape.text_frame
        tf_c.word_wrap = True
        tf_c.vertical_anchor = MSO_ANCHOR.TOP
        tf_c.margin_top = Inches(0.04)
        tf_c.margin_left = Inches(0.08)
        tf_c.margin_right = Inches(0.08)
        
        p0 = tf_c.paragraphs[0]
        p0.text = title_c
        p0.font.name = "Arial"
        p0.font.size = Pt(7.5)
        p0.font.bold = True
        p0.font.color.rgb = RGBColor(106, 27, 154)
        p0.space_after = Pt(1)
        
        p1 = tf_c.add_paragraph()
        p1.text = desc_c
        p1.font.name = "Arial"
        p1.font.size = Pt(6.5)
        p1.font.color.rgb = RGBColor(71, 85, 105)

    # Arrow 3: Data Exchange
    arr_z3 = slide.shapes.add_shape(MSO_SHAPE.LEFT_RIGHT_ARROW, Inches(7.44), arch_top + Inches(2.45), Inches(0.20), Inches(0.14))
    arr_z3.fill.solid()
    arr_z3.fill.fore_color.rgb = RGBColor(100, 116, 139)
    arr_z3.line.color.rgb = RGBColor(100, 116, 139)

    # -------------------------------------------------------------
    # ZONE 4: Data & Intelligence Layer (The "Brain")
    # -------------------------------------------------------------
    z4_x = Inches(7.68)
    z4_w = Inches(2.15)
    
    z4_hdr = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, z4_x, arch_top + Inches(0.08), z4_w, Inches(0.40))
    z4_hdr.fill.solid()
    z4_hdr.fill.fore_color.rgb = RGBColor(224, 247, 250) # Cyan
    z4_hdr.line.color.rgb = RGBColor(128, 222, 234)
    tf4 = z4_hdr.text_frame
    tf4.word_wrap = True
    tf4.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf4.paragraphs[0]
    p.text = "Zone 4: AI & ML Engine"
    p.font.name = "Arial"
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 131, 143)
    p.alignment = PP_ALIGN.CENTER
    
    z4_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, z4_x, arch_top + Inches(0.56), z4_w, Inches(3.86))
    z4_box.fill.solid()
    z4_box.fill.fore_color.rgb = RGBColor(240, 251, 252)
    z4_box.line.color.rgb = RGBColor(128, 222, 234)
    
    z4_cards = [
        ("Ensemble ML Predictor", "LightGBM + XGBoost 2.1 regressors predict residual delay days."),
        ("TreeSHAP Attribution", "Micro-level factor attribution quantifying additive day delays."),
        ("Kaplan-Meier Survival", "Probabilistic time-to-lapse modeling for 365-day Sec 3D deadlines."),
        ("Prescriptive Rule Core", "Statutory threshold trigger engine activating specific SOP orders."),
        ("Central DB (<15ms)", "High-speed normalized SQLite/PostgreSQL storage across 47 attributes.")
    ]
    
    for i, (title_c, desc_c) in enumerate(z4_cards):
        cy = arch_top + Inches(0.62) + i * Inches(0.75)
        c_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, z4_x + Inches(0.06), cy, z4_w - Inches(0.12), card_h)
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
        c_shape.line.color.rgb = RGBColor(178, 235, 242)
        tf_c = c_shape.text_frame
        tf_c.word_wrap = True
        tf_c.vertical_anchor = MSO_ANCHOR.TOP
        tf_c.margin_top = Inches(0.04)
        tf_c.margin_left = Inches(0.08)
        tf_c.margin_right = Inches(0.08)
        
        p0 = tf_c.paragraphs[0]
        p0.text = title_c
        p0.font.name = "Arial"
        p0.font.size = Pt(7.5)
        p0.font.bold = True
        p0.font.color.rgb = RGBColor(0, 131, 143)
        p0.space_after = Pt(1)
        
        p1 = tf_c.add_paragraph()
        p1.text = desc_c
        p1.font.name = "Arial"
        p1.font.size = Pt(6.5)
        p1.font.color.rgb = RGBColor(71, 85, 105)

    # =============================================================
    # 2. BOTTOM BAR: COMPONENTS / TECHNOLOGY STACK (REAL LOGOS!)
    # =============================================================
    bot_y = Inches(5.78)
    bot_h = Inches(0.96)
    
    bot_bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.45), bot_y, Inches(9.45), bot_h)
    bot_bg.fill.solid()
    bot_bg.fill.fore_color.rgb = RGBColor(255, 255, 255)
    bot_bg.line.color.rgb = RGBColor(203, 213, 225)
    bot_bg.line.width = Pt(1.2)
    
    # Left Header Label
    lbl_tech = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.52), bot_y + Inches(0.10), Inches(1.50), bot_h - Inches(0.20))
    lbl_tech.fill.solid()
    lbl_tech.fill.fore_color.rgb = RGBColor(241, 245, 249)
    lbl_tech.line.color.rgb = RGBColor(203, 213, 225)
    tf_lt = lbl_tech.text_frame
    tf_lt.word_wrap = True
    tf_lt.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf_lt.paragraphs[0]
    p.text = "Core Technology\nStack & Engine"
    p.font.name = "Arial"
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(15, 23, 42)
    p.alignment = PP_ALIGN.CENTER
    
    # 9 Authentic Tech Badges with REAL IMAGE LOGOS!
    tech_logos = [
        ("python_logo.png", "Python 3.13", "FastAPI Core", RGBColor(56, 126, 184)),
        ("fastapi_logo.png", "FastAPI", "REST Gateway", RGBColor(0, 150, 136)),
        ("react_logo.png", "React 19", "Vite Frontend", RGBColor(32, 35, 42)),
        ("xgboost_logo.png", "XGBoost 2.1", "ML Regressor", RGBColor(230, 81, 0)),
        ("lightgbm_logo.png", "LightGBM", "Ensemble Trees", RGBColor(13, 71, 161)),
        ("shap_logo.png", "TreeSHAP", "Statutory XAI", RGBColor(30, 41, 59)),
        ("leaflet_logo.png", "Leaflet GIS", "EPSG:4326 Maps", RGBColor(116, 172, 0)),
        ("postgresql_logo.png", "PostgreSQL", "<15ms DB Cache", RGBColor(51, 103, 145)),
        ("sha256_logo.png", "SHA-256", "Audit Ledger", RGBColor(79, 70, 229))
    ]
    
    b_w = Inches(0.81)
    b_gap = Inches(0.06)
    start_bx = Inches(2.08)
    
    for i, (img_name, b_name, b_tag, b_border) in enumerate(tech_logos):
        bx = start_bx + i * (b_w + b_gap)
        
        # Badge card background
        card_b = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, bot_y + Inches(0.08), b_w, bot_h - Inches(0.16))
        card_b.fill.solid()
        card_b.fill.fore_color.rgb = RGBColor(248, 250, 252)
        card_b.line.color.rgb = b_border
        card_b.line.width = Pt(1.2)
        
        # Embed Real Image Logo!
        img_path = os.path.join('sih/assets', img_name)
        slide.shapes.add_picture(img_path, bx + Inches(0.24), bot_y + Inches(0.12), width=Inches(0.33), height=Inches(0.33))
        
        # Text under logo
        tf_b = card_b.text_frame
        tf_b.word_wrap = True
        tf_b.vertical_anchor = MSO_ANCHOR.BOTTOM
        tf_b.margin_left = Inches(0.02)
        tf_b.margin_right = Inches(0.02)
        tf_b.margin_bottom = Inches(0.04)
        
        p0 = tf_b.paragraphs[0]
        p0.text = b_name
        p0.font.name = "Arial"
        p0.font.size = Pt(6.5)
        p0.font.bold = True
        p0.font.color.rgb = RGBColor(15, 23, 42)
        p0.alignment = PP_ALIGN.CENTER
        
        p1 = tf_b.add_paragraph()
        p1.text = b_tag
        p1.font.name = "Arial"
        p1.font.size = Pt(5.5)
        p1.font.color.rgb = RGBColor(100, 116, 139)
        p1.alignment = PP_ALIGN.CENTER

    # =============================================================
    # 3. RIGHT COLUMN: IMPLEMENTATION PROCESS (REAL STEP ICONS!)
    # =============================================================
    imp_x = Inches(10.05)
    imp_w = Inches(2.80)
    imp_top = Inches(1.12)
    imp_h = Inches(5.62)
    
    # Outer Container Box
    imp_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, imp_x, imp_top, imp_w, imp_h)
    imp_box.fill.solid()
    imp_box.fill.fore_color.rgb = RGBColor(255, 255, 255)
    imp_box.line.color.rgb = RGBColor(203, 213, 225)
    imp_box.line.width = Pt(1.2)
    
    # Header: Implementation Process (Crimson in SIHex1!)
    imp_hdr = slide.shapes.add_textbox(imp_x, imp_top + Inches(0.06), imp_w, Inches(0.35))
    tf_ih = imp_hdr.text_frame
    p = tf_ih.paragraphs[0]
    p.text = "Implementation Process"
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RGBColor(194, 24, 91) # Crimson / Magenta
    p.alignment = PP_ALIGN.CENTER
    
    # Vertical Connecting Spine
    spine = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, imp_x + Inches(0.38), imp_top + Inches(0.60), Inches(0.04), Inches(4.70))
    spine.fill.solid()
    spine.fill.fore_color.rgb = RGBColor(226, 232, 240)
    spine.line.color.rgb = RGBColor(203, 213, 225)
    spine.line.width = Pt(0.5)
    
    # 6 Process Steps with Real High-Resolution Step Icons!
    steps_data = [
        ("step1_ingest.png", "1. Multi-Modal Ingestion", 
         "Gazette PDFs, e-Courts stays, cadastral Jamabandi records & BhoomiRashi data parsed into pipeline.", 
         RGBColor(0, 105, 92)),
        ("step2_centraldb.png", "2. Decoupled Central DB", 
         "Normalizes 47 statutory attributes into high-speed SQLite/PostgreSQL cache with <15ms queries.", 
         RGBColor(38, 50, 56)),
        ("step3_ml.png", "3. Ensemble ML Inference", 
         "XGBoost 2.1 & LightGBM regressors predict residual acquisition days & statutory lapse risks.", 
         RGBColor(230, 81, 0)),
        ("step4_shap.png", "4. TreeSHAP Attribution", 
         "Deconstructs prediction into exact additive marginal delay days per statutory obstacle.", 
         RGBColor(51, 105, 30)),
        ("step5_dfa.png", "5. Prescriptive Legal DFA", 
         "Automatically drafts statutory Section 3D, 3G & 3H compliance orders for nodal officers.", 
         RGBColor(191, 54, 12)),
        ("step6_esign.png", "6. HITL e-Sign & Audit", 
         "Nodal officer validates order via Aadhaar DSC; seals tamper-proof SHA-256 block ledger for CAG.", 
         RGBColor(136, 14, 79))
    ]
    
    step_start_y = imp_top + Inches(0.48)
    step_spacing = Inches(0.83)
    
    for i, (icon_file, s_title, s_desc, s_color) in enumerate(steps_data):
        sy = step_start_y + i * step_spacing
        
        # Real Custom Step Icon Graphic Image!
        icon_path = os.path.join('sih/assets', icon_file)
        slide.shapes.add_picture(icon_path, imp_x + Inches(0.16), sy + Inches(0.08), width=Inches(0.46), height=Inches(0.46))
        
        # Step Text Box
        txt_box = slide.shapes.add_textbox(imp_x + Inches(0.68), sy, imp_w - Inches(0.74), Inches(0.80))
        tf_txt = txt_box.text_frame
        tf_txt.word_wrap = True
        tf_txt.vertical_anchor = MSO_ANCHOR.TOP
        tf_txt.margin_left = Inches(0.02)
        tf_txt.margin_right = Inches(0.02)
        tf_txt.margin_top = Inches(0.04)
        tf_txt.margin_bottom = Inches(0.02)
        
        p_t = tf_txt.paragraphs[0]
        p_t.text = s_title
        p_t.font.name = "Arial"
        p_t.font.size = Pt(8.5)
        p_t.font.bold = True
        p_t.font.color.rgb = s_color
        p_t.space_after = Pt(1)
        
        p_d = tf_txt.add_paragraph()
        p_d.text = s_desc
        p_d.font.name = "Arial"
        p_d.font.size = Pt(7)
        p_d.font.color.rgb = RGBColor(71, 85, 105)

    prs.save('sih/SIH-2026_template_backup.pptx')
    prs.save('sih/GatiShakti_AI_SIH2026_Presentation.pptx')
    print("Slide 3 with real assets & logos built successfully in both PPTX files!")

if __name__ == '__main__':
    build_slide_3()
